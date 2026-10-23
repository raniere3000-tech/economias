"""Gera uma lista de comparação "Virar para Água e Esgoto" a partir da base sintética, para testar o comparativo.
Uso: python tests/gerar_lista_teste.py base_teste.xlsx AAAA-MM "Julho - Virar para Água e Esgoto.xlsx"
Pega matrículas que começam a faturar esgoto no mês, antes, depois, que já faturavam em dez/25 e que nunca faturam,
mais 15 matrículas que não existem na base. O nome do arquivo precisa ter "Virar" e o nome do mês."""
import sys, random, datetime
from openpyxl import load_workbook, Workbook
base, mes, saida = sys.argv[1], sys.argv[2], sys.argv[3]
ano, mm = map(int, mes.split('-'))
random.seed(mm)
pri, cad = {}, {}
for ws in load_workbook(base, read_only=True).worksheets:
    cab = None
    for row in ws.iter_rows(values_only=True):
        if cab is None:
            if row and 'NUM_LIGACAO' in row: cab = {c: i for i, c in enumerate(row) if c}
            continue
        g = lambda n: row[cab[n]]
        if str(g('ULTIMO_CONTRATO')).upper() != 'SIM': continue
        m, d = g('NUM_LIGACAO'), g('Referência'); k = (d.year, d.month)
        cad[m] = row
        if (g('Qtd_eco_fat_esgoto') or 0) > 0 and (m not in pri or k < pri[m]): pri[m] = k
alvo = (ano, mm)
grupos = {'no mes': [m for m in cad if pri.get(m) == alvo], 'antes': [m for m in cad if m in pri and (2026, 1) <= pri[m] < alvo],
          'depois': [m for m in cad if m in pri and pri[m] > alvo], 'dez': [m for m in cad if pri.get(m) == (2025, 12)],
          'nunca': [m for m in cad if m not in pri]}
qtd = {'no mes': 0.6, 'antes': 0.04, 'depois': 0.05, 'dez': 0.1, 'nunca': 0.02}
escolhidas = []
for g, ms in grupos.items():
    escolhidas += random.sample(ms, int(len(ms) * qtd[g]))
wb = Workbook(write_only=True); ws = wb.create_sheet(saida.rsplit('.', 1)[0][:31])
hdr = ['NUM_LIGACAO','NOM_CLIENTE','CIDADE','SUB_CATEGORIA','SIT_LIG','SIT_CONTRATO','Referência','ULTIMO_CONTRATO','COD_GRUPO','TIPO_FATURAMENTO','Faturamento_bruto_direta_agua','Faturamento_bruto_direta_esgoto','Qtd_eco_fat_agua','Qtd_eco_fat_esgoto','Eco_total']
ws.append(hdr)
ref = datetime.datetime(ano, mm, 1)
for m in sorted(escolhidas):
    r = list(cad[m]); r[6] = ref; r[13] = r[14]; ws.append(r)
for i in range(15):
    ws.append([900000000 + mm * 100 + i, f'FORA DA BASE {i}', 'MACAÉ', 'RESIDENCIAL', 'ATIVA', 'ATIVO', ref, 'SIM', 511, 'AGUA E ESGOTO', 50, 40, 1, 1, 1])
wb.save(saida)
print('ok', {g: int(len(ms) * qtd[g]) for g, ms in grupos.items()}, '+15 fora da base')
