"""Gera um cadastro de dez/25 sintético, no formato do cadastro real, a partir da base sintética do estudo.
Uso: python tests/gerar_cadastro_teste.py base_teste.xlsx "Base_Cadastro_Dezembro-2025.xlsx"
Cobre todos os casos das regras: ativas, inativas, desativadas, cortadas, corte pedido e provisórias; matrículas que
aparecem só em 2026 com e sem cadastro; faturando em dez/25 fora do cadastro e inativas faturando (conferências);
TIPO_FATURAMENTO AGUA e AGUA E ESGOTO; economias do cadastro diferentes das faturadas; e matrículas só do cadastro."""
import sys, random, collections
from openpyxl import Workbook, load_workbook

base, saida = sys.argv[1], sys.argv[2]
rnd = random.Random(23)
info = {}; dez = None
for ws in load_workbook(base, read_only=True).worksheets:
    cab = None
    for row in ws.iter_rows(values_only=True):
        if cab is None:
            if row and 'NUM_LIGACAO' in row: cab = {c: i for i, c in enumerate(row) if c}
            continue
        g = lambda n: row[cab[n]]
        k = (g('Referência').year, g('Referência').month); dez = min(dez or k, k)
        m = g('NUM_LIGACAO'); d = info.setdefault(m, {'cad': (g('NOM_CLIENTE'), g('CIDADE'), g('SUB_CATEGORIA')), 'mes': {}})
        q = d['mes'].get(k, (0, 0, 0))
        d['mes'][k] = (max(q[0], int(g('Qtd_eco_fat_agua') or 0)), max(q[1], int(g('Qtd_eco_fat_esgoto') or 0)), max(q[2], int(g('Eco_total') or 0)))
wb = Workbook(); ws = wb.active; ws.title = 'Query1'
ws.append(['NUM_LIGACAO', 'NOM_CLIENTE', 'CIDADE', 'SIT_CONTRATO', 'SIT_LIG', 'SUB_CATEGORIA', 'QTD_ECO_COM', 'QTD_ECO_IND', 'QTD_ECO_PUB', 'QTD_ECO_RES', 'QTD_ECO_OUT', 'TOTAL_ECO', 'TIPO_FATURAMENTO'])
cont = collections.Counter()
def linha(m, nome, cid, sub, sit, eco, tipo):
    contr = {'ATIVA': 'ATIVO', 'CORTADA': 'ATIVO', 'CORTE PEDIDO': 'ATIVO', 'PROVISÓRIA': 'PENDENTE'}.get(sit, 'ENCERRADO')
    res = max(eco - 1, 0) if eco > 1 else eco; com = eco - res
    ws.append([m, nome, cid, contr, sit, sub, com, 0, 0, res, 1 if rnd.random() < 0.05 else 0, eco, tipo]); cont[sit] += 1
for i, (m, d) in enumerate(sorted(info.items(), key=lambda x: str(x[0]))):
    nome, cid, sub = d['cad']; qa, qe, et = d['mes'].get(dez, (0, 0, 0)); na_dez = dez in d['mes']
    eco = et or max(x[2] for x in d['mes'].values()) or 1
    if rnd.random() < 0.15: eco = max(1, eco + rnd.choice([-1, 1]))          # economias do cadastro diferentes
    if na_dez and (qa or qe):
        if i % 53 == 0: continue                                             # faturando fora do cadastro
        sit = 'INATIVA' if i % 47 == 0 else 'CORTADA' if i % 31 == 0 else 'ATIVA'
        tipo = 'AGUA E ESGOTO' if qe else 'AGUA'
    elif na_dez:
        sit = ['ATIVA', 'INATIVA', 'DESATIVADA', 'CORTADA', 'CORTE PEDIDO'][i % 5]
        tipo = rnd.choice(['AGUA', 'AGUA', 'AGUA E ESGOTO'])
    else:                                                                    # aparece só em 2026
        x = rnd.random()
        if x < 0.4: continue                                                 # fora do cadastro: nova na base
        sit = 'PROVISÓRIA' if x < 0.55 else rnd.choice(['INATIVA', 'DESATIVADA', 'CORTADA', 'ATIVA'])
        tipo = rnd.choice(['AGUA', 'AGUA E ESGOTO'])
    linha(m, nome, cid, sub, sit, eco, tipo)
for i in range(60):                                                          # só no cadastro (ativas sem faturar fora da base)
    linha(900000 + i, f'CLIENTE SO CADASTRO {i}', 'CIDADE A', 'RESIDENCIAL', 'ATIVA', rnd.randint(1, 3), rnd.choice(['AGUA', 'AGUA E ESGOTO']))
wb.save(saida)
print('cadastro:', sum(cont.values()), 'matrículas', dict(cont))
