"""Recalcula em Python, de forma independente, os números do Resumo do painel.
Uso: python tests/conferir_totais.py base_teste.xlsx
Compare a saída com os indicadores e as tabelas do painel (mesmas regras do index.html)."""
import sys, collections
from openpyxl import load_workbook

arq = sys.argv[1] if len(sys.argv) > 1 else 'base_teste.xlsx'
wb = load_workbook(arq, read_only=True)
T, Q, F = (collections.defaultdict(dict) for _ in range(3))
refs = set()
for ws in wb.worksheets:
    cab = None
    for row in ws.iter_rows(values_only=True):
        if cab is None:
            if row and 'NUM_LIGACAO' in [str(c).strip().upper() if c else '' for c in row]:
                cab = {str(c).strip(): i for i, c in enumerate(row) if c}
            continue
        g = lambda nome: row[cab[nome]]
        if str(g('ULTIMO_CONTRATO') or '').strip().upper() != 'SIM':
            continue
        d = g('Referência'); k = (d.year, d.month); refs.add(k); m = g('NUM_LIGACAO')
        T[m][k] = 'ESGOTO' in str(g('TIPO_FATURAMENTO') or '').upper()
        Q[m][k] = int(g('Qtd_eco_fat_esgoto') or 0)
        F[m][k] = float(g('Faturamento_bruto_direta_esgoto') or 0)
refs = sorted(refs); n = len(refs)
z = lambda: [0.0] * n
nov, aum, sai, red, ree, fat_tot, fat_novas = z(), z(), z(), z(), z(), z(), z()
e_base, e_nova, lig, lig_nova, lig_aum = z(), z(), z(), z(), z()
for m in T:
    h = [T[m].get(k, False) for k in refs]; q = [Q[m].get(k, 0) for k in refs]; f = [F[m].get(k, 0) for k in refs]
    entra = -1
    if not h[0]:
        for r in range(1, n):
            if h[r] and not h[r - 1]: entra = r; break
    # matrícula nova: aparece na base pela primeira vez em 2026 já com AGUA E ESGOTO
    prim = min(refs.index(k) for k in T[m])
    nova = entra > 0 and entra == prim
    e_orig = e_nova if nova else e_base
    for r in range(n):
        fat_tot[r] += f[r]
        if r == 0: continue
        if h[r]:
            if entra > 0 and r >= entra: fat_novas[r] += f[r]
            elif h[0] and q[r] > q[0] and q[r] > 0: fat_novas[r] += f[r] * (q[r] - q[0]) / q[r]
        if h[r] and not h[r - 1]:
            if r == entra:
                nov[r] += q[r]; e_orig[r] += q[r]; lig[r] += 1
                if nova: lig_nova[r] += 1
            else: ree[r] += q[r]
        elif h[r - 1] and not h[r]: sai[r] += q[r - 1]
        elif h[r] and h[r - 1]:
            d = q[r] - q[r - 1]
            if d > 0: aum[r] += d; e_orig[r] += d; lig_aum[r] += 1
            elif d < 0: red[r] -= d
rot = [f'{["jan","fev","mar","abr","mai","jun","jul","ago","set","out","nov","dez"][mm-1]}/{str(a)[2:]}' for a, mm in refs]
print('mês      na base  mat.novas  novas  ligações  saídas  reentr.  fat_novas      fat_total')
for r in range(n):
    print(f'{rot[r]:7} {e_base[r]:8.0f} {e_nova[r]:10.0f} {nov[r]+aum[r]:6.0f} {lig[r]:9.0f} {sai[r]+red[r]:7.0f} {ree[r]:8.0f} {fat_novas[r]:11,.0f} {fat_tot[r]:14,.0f}')
print(f'Economias novas em 2026: {sum(nov)+sum(aum):,.0f}  (já na base {sum(e_base):,.0f} + matrículas novas {sum(e_nova):,.0f})')
print(f'Ligações que passaram a faturar esgoto em 2026: {sum(lig):,.0f}  (já na base {sum(lig)-sum(lig_nova):,.0f} + matrículas novas {sum(lig_nova):,.0f})')
print('Por mês, ligações que já tinham esgoto e aumentaram:', ' '.join(f'{rot[r]} {lig_aum[r]:.0f}' for r in range(1, n)))
print('Por mês, economias acrescentadas nelas:', ' '.join(f'{rot[r]} {aum[r]:.0f}' for r in range(1, n)))
print(f'Fat. esgoto das novas no ano: {sum(fat_novas):,.0f}   Saídas: {sum(sai)+sum(red):,.0f}   Reentradas: {sum(ree):,.0f}')
