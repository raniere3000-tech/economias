"""Lê a pasta de arquivos do painel de faturamento (fatura + consumo por mês) e confere a aba "Arquivos do faturamento".
Uso: python tests/conferir_faturamento.py PASTA [--base saida.xlsx]
- Fatura: colunas Rubrica, Valor Parcela e Data de Vencimento (arquivo com "avulso" no nome fica de fora); valem as
  rubricas VALOR DE AGUA e VALOR DE ESGOTO; o mês é a Referencia de Leitura.
- Consumo: colunas Leitura Atual, Consumo Medido e Consumo Faturado; o mês vem do nome do arquivo (MM-AAAA ou mes-AA),
  do nome da pasta ou de Mes/Ano Lancamento.
- Economias faturadas de água (esgoto) da ligação no mês: soma de Qtd. Economia Residencial, Comercial, Industrial e
  Publica do consumo, quando o Consumo Faturado é maior que 0 e há linha de água (esgoto) na fatura. Uma vez por
  ligação; o painel de faturamento soma por linha de fatura (também mostrado aqui).
--base: grava os dados no formato da base do estudo, para conferir o painel com tests/conferir_totais.py."""
import sys, os, re, csv, collections, datetime, unicodedata
from openpyxl import load_workbook, Workbook

pasta = sys.argv[1]
saida = sys.argv[sys.argv.index('--base') + 1] if '--base' in sys.argv else None
LIG = ['N. Ligacao', 'N. da Ligacao', 'N Ligacao', 'N. da Ligação', 'N. Ligação']
ECO = ['Qtd. Economia Residencial', 'Qtd. Economia Comercial', 'Qtd. Economia Industrial', 'Qtd. Economia Publica']
MES = ['jan', 'fev', 'mar', 'abr', 'mai', 'jun', 'jul', 'ago', 'set', 'out', 'nov', 'dez']
sem_acento = lambda s: unicodedata.normalize('NFD', str(s)).encode('ascii', 'ignore').decode()
limpa = lambda c: str(c or '').strip().replace('﻿', '').replace('"', '').replace('Pública', 'Publica')

def linhas(caminho):
    if caminho.lower().endswith('.csv'):
        raw = open(caminho, 'rb').read()
        try: txt = raw.decode('utf-8-sig')
        except UnicodeDecodeError: txt = raw.decode('latin-1')
        yield from csv.reader(txt.splitlines(), delimiter=';')
    else:
        for ws in load_workbook(caminho, read_only=True).worksheets:
            for r in ws.iter_rows(values_only=True): yield ['' if v is None else v for v in r]

def num(v):
    if isinstance(v, (int, float)): return float(v)
    s = str(v or '').strip()
    if not s: return 0.0
    if ',' in s: s = s.replace('.', '').replace(',', '.')
    try: return float(s)
    except ValueError: return 0.0

def ref_texto(v):
    if isinstance(v, datetime.datetime): return f'{v.year}-{v.month:02d}'
    s = sem_acento(str(v or '')).strip().lower()
    m = re.match(r'^(\d{1,2})/(\d{1,2})/(\d{4})', s) or None
    if m: return f'{m.group(3)}-{int(m.group(2)):02d}'
    m = re.match(r'^(\d{1,2})[/-](\d{4})$', s)
    if m: return f'{m.group(2)}-{int(m.group(1)):02d}'
    m = re.match(r'^(\d{4})-(\d{2})', s)
    if m: return f'{m.group(1)}-{m.group(2)}'
    m = re.match(r'^([a-z]{3})[a-z]*[/\-. ](\d{2}|\d{4})$', s)
    if m and m.group(1) in MES: return f'{m.group(2) if len(m.group(2)) == 4 else "20" + m.group(2)}-{MES.index(m.group(1)) + 1:02d}'
    return None

def ref_nome(nome):
    n = sem_acento(nome).lower()
    m = re.search(r'(?<!\d)(\d{2})[-/._ ](\d{4})(?!\d)', n)
    if m and 1 <= int(m.group(1)) <= 12: return f'{m.group(2)}-{m.group(1)}'
    m = re.search(r'([a-z]{3})[a-z]*[-/._ ](\d{4}|\d{2})(?!\d)', n)
    if m and m.group(1) in MES: return f'{m.group(2) if len(m.group(2)) == 4 else "20" + m.group(2)}-{MES.index(m.group(1)) + 1:02d}'
    return None

fat = collections.defaultdict(lambda: {'fa': 0.0, 'fe': 0.0, 'la': 0, 'le': 0})
cad = {}
cons = {}
for raiz, _, nomes in os.walk(pasta):
    for nome in sorted(nomes):
        if not re.search(r'\.(csv|xlsx)$', nome, re.I) or nome.startswith('~$'): continue
        caminho = os.path.join(raiz, nome); it = linhas(caminho); cab = None
        for i, row in enumerate(it):
            cols = [limpa(c) for c in row]
            if any(c in LIG for c in cols): cab = {c: k for k, c in enumerate(cols) if c}; break
            if i > 30: break
        if not cab: continue
        g = lambda r, c: r[cab[c]] if c in cab and cab[c] < len(r) else ''
        cl = next(c for c in LIG if c in cab)
        if {'Rubrica', 'Valor Parcela', 'Data de Vencimento'} <= set(cab) and 'avulso' not in nome.lower():
            for r in it:
                rub = str(g(r, 'Rubrica')).strip().upper()
                if rub not in ('VALOR DE AGUA', 'VALOR DE ESGOTO'): continue
                k = ref_texto(g(r, 'Referencia de Leitura'))
                if not k: continue
                m = str(g(r, cl)).strip()
                if m.endswith('.0'): m = m[:-2]
                x = fat[(m, k)]
                if rub == 'VALOR DE AGUA': x['fa'] += num(g(r, 'Valor Parcela')); x['la'] += 1
                else: x['fe'] += num(g(r, 'Valor Parcela')); x['le'] += 1
                if m not in cad or k >= cad[m][0]:
                    cad[m] = (k, g(r, 'Nome Cliente'), g(r, 'Nome da Localidade'), g(r, 'Categoria'), g(r, 'Situacao Ligacao'), g(r, 'Situacao Conta'), g(r, 'Grupo'))
        elif {'Leitura Atual', 'Consumo Medido', 'Consumo Faturado'} <= set(cab):
            k = ref_nome(nome) or ref_nome(os.path.basename(raiz) + ' ')
            for r in it:
                kk = k or (f"{int(num(g(r, 'Ano Lancamento')))}-{int(num(g(r, 'Mes Lancamento'))):02d}" if g(r, 'Mes Lancamento') else None)
                if not kk: continue
                m = str(g(r, cl)).strip()
                if m.endswith('.0'): m = m[:-2]
                cons[(m, kk)] = (num(g(r, 'Consumo Faturado')), sum(num(g(r, c)) for c in ECO))
chaves = set(fat) | set(cons)
refs = sorted({k for _, k in chaves})
tot = {k: collections.Counter() for k in refs}
base = []
for (m, k) in sorted(chaves):
    x = fat.get((m, k), {'fa': 0.0, 'fe': 0.0, 'la': 0, 'le': 0}); c, eco = cons.get((m, k), (0.0, 0))
    qa = eco if c > 0 and x['la'] else 0; qe = eco if c > 0 and x['le'] else 0
    t = tot[k]; t['fa'] += x['fa']; t['fe'] += x['fe']; t['qa'] += qa; t['qe'] += qe
    t['qa_lin'] += eco * x['la'] if c > 0 else 0; t['qe_lin'] += eco * x['le'] if c > 0 else 0
    base.append((m, k, x['fa'], x['fe'], qa, qe, eco))
print('mês       R$ água        R$ esgoto     eco água  (por linha)   eco esgoto  (por linha)')
for k in refs:
    t = tot[k]; print(f'{k}  {t["fa"]:13,.2f}  {t["fe"]:13,.2f}  {t["qa"]:9,.0f}  {t["qa_lin"]:11,.0f}  {t["qe"]:11,.0f}  {t["qe_lin"]:11,.0f}')
if saida:
    wb = Workbook(write_only=True); ws = wb.create_sheet('Base')
    ws.append(['NUM_LIGACAO','NOM_CLIENTE','CIDADE','SUB_CATEGORIA','SIT_LIG','SIT_CONTRATO','Referência','ULTIMO_CONTRATO','COD_GRUPO','TIPO_FATURAMENTO','Faturamento_bruto_direta_agua','Faturamento_bruto_direta_esgoto','Qtd_eco_fat_agua','Qtd_eco_fat_esgoto','Eco_total'])
    for m, k, fa, fe, qa, qe, eco in base:
        _, nome, cid, cat, sl, sc, grp = cad.get(m, ('', '', '', '', '', '', ''))
        ws.append([m, nome, cid, cat, sl, sc, datetime.datetime(int(k[:4]), int(k[5:]), 1), 'SIM', grp, '', round(fa, 2), round(fe, 2), int(qa), int(qe), int(eco)])
    wb.save(saida); print('base gravada:', saida, len(base), 'linhas')
